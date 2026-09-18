from __future__ import annotations

import base64
import ctypes
import json
import re
import sys
import threading
import time
from dataclasses import dataclass, asdict
from datetime import datetime
from pathlib import Path
from typing import Any
from urllib.parse import urljoin

# 使用 Windows 憑證存放區，讓企業/系統 CA 能被 HTTPS 驗證使用。
import truststore
truststore.inject_into_ssl()

import requests
from requests.adapters import HTTPAdapter
from urllib3.util import Retry
from bs4 import BeautifulSoup
from google import genai
from google.genai import types
from PySide6.QtCore import QSettings, QThread, Qt, Signal, QUrl, QStandardPaths
from PySide6.QtGui import QDesktopServices
from PySide6.QtWidgets import (
    QApplication,
    QCheckBox,
    QComboBox,
    QFormLayout,
    QFileDialog,
    QGridLayout,
    QGroupBox,
    QHBoxLayout,
    QHeaderView,
    QLabel,
    QLineEdit,
    QListWidget,
    QMainWindow,
    QMessageBox,
    QPushButton,
    QSpinBox,
    QSplitter,
    QTableWidget,
    QTableWidgetItem,
    QVBoxLayout,
    QWidget,
)


PTT_BASE = "https://www.ptt.cc"
DEFAULT_BOARDS = ["MacShop", "mobilesales", "iOS", "MobileComm", "nb-shopping", "HardwareSale"]
GEMINI_MODEL = "gemini-2.5-flash"
SEARCH_MODE_PYTHON = "python"
SEARCH_MODE_GEMINI = "gemini"
SEARCH_MODES = [
    ("純 Python 關鍵字搜尋（推薦，不需 API Key）", SEARCH_MODE_PYTHON),
    ("Gemini AI 結構化分析", SEARCH_MODE_GEMINI),
]
COLOR_TERMS_CHINESE = [
    "星光色", "太空灰", "午夜色", "原色鈦金屬", "沙漠色", "自然鈦", "黑鈦", "白鈦", "藍鈦",
    "黑色", "白色", "金色", "銀色", "銀", "藍色", "紫色", "粉色", "紅色", "綠色", "黃色", "灰色", "橘色", "橙色",
]
COLOR_TERMS_ENGLISH = [
    "black", "white", "gold", "silver", "blue", "purple", "pink", "red", "green", "yellow", "gray", "grey",
]
CAPACITY_PATTERN = re.compile(
    r"(?<![A-Za-z0-9])(32|64|128|256|512|1024|[12])\s*(GB|G|TB|T)(?![A-Za-z0-9])", re.I
)
PRICE_PATTERN = re.compile(
    r"(?:(?:售價|價格|售|賣價|售出價|收)?\s*[:：]?\s*)?"
    r"(?:NT\$|NTD\s*)?([0-9]{1,3}(?:,[0-9]{3})+|[0-9]{4,6})\s*(?:元|塊|新台幣|TWD)?",
    re.I,
)
CHINESE_PRICE_PATTERN = re.compile(r"(\d{1,2})\s*萬\s*(\d{0,4})\s*(?:元|塊|新台幣)?")
SOLD_TERMS = ("已售出", "已賣出", "已成交", "交易完成", "已預訂", "已預定", "sold", "reserved")
UNSOLD_TERMS = ("未售出", "尚未售出", "還沒售出", "未賣出")
GEMINI_MODELS = [
    ("gemini-2.5-flash（推薦，快速低成本）", "gemini-2.5-flash"),
    ("gemini-2.5-pro（高品質推理）", "gemini-2.5-pro"),
    ("gemini-2.0-flash", "gemini-2.0-flash"),
    ("gemini-1.5-flash", "gemini-1.5-flash"),
    ("gemini-1.5-pro", "gemini-1.5-pro"),
]


class _DataBlob(ctypes.Structure):
    _fields_ = [("cbData", ctypes.c_uint32), ("pbData", ctypes.POINTER(ctypes.c_ubyte))]


def _crypt_with_windows_dpapi(value: bytes, protect: bool, flags: int = 0) -> bytes:
    """用 Windows DPAPI 對目前使用者加解密祕密資料。"""
    if sys.platform != "win32":
        raise OSError("Windows DPAPI 僅適用於 Windows")

    crypt32 = ctypes.WinDLL("crypt32", use_last_error=True)
    kernel32 = ctypes.WinDLL("kernel32", use_last_error=True)
    kernel32.LocalFree.argtypes = [ctypes.c_void_p]
    kernel32.LocalFree.restype = ctypes.c_void_p
    input_buffer = ctypes.create_string_buffer(value)
    input_blob = _DataBlob(
        len(value), ctypes.cast(input_buffer, ctypes.POINTER(ctypes.c_ubyte))
    )
    output_blob = _DataBlob()
    function = crypt32.CryptProtectData if protect else crypt32.CryptUnprotectData
    function.argtypes = [
        ctypes.POINTER(_DataBlob),
        ctypes.c_wchar_p,
        ctypes.POINTER(_DataBlob),
        ctypes.c_void_p,
        ctypes.c_void_p,
        ctypes.c_uint32,
        ctypes.POINTER(_DataBlob),
    ]
    function.restype = ctypes.c_bool
    if not function(
        ctypes.byref(input_blob), None, None, None, None, flags, ctypes.byref(output_blob)
    ):
        raise ctypes.WinError(ctypes.get_last_error())
    try:
        return ctypes.string_at(output_blob.pbData, output_blob.cbData)
    finally:
        kernel32.LocalFree(output_blob.pbData)


def protect_api_key(api_key: str) -> str:
    """回傳可放入 QSettings 的加密 API Key；失敗時不保存明文。"""
    raw = api_key.encode("utf-8")
    try:
        encrypted = _crypt_with_windows_dpapi(raw, protect=True)
        scope = "user"
    except OSError:
        # 精簡/隔離 Windows 帳號可能沒有可用的使用者 DPAPI profile；
        # 機器範圍仍是加密保存，並以前綴標示解密方式。
        encrypted = _crypt_with_windows_dpapi(raw, protect=True, flags=0x4)
        scope = "machine"
    return f"{scope}:{base64.b64encode(encrypted).decode('ascii')}"


def unprotect_api_key(encoded: str) -> str:
    scope, separator, payload = encoded.partition(":")
    if not separator or scope not in {"user", "machine"}:
        # 相容早期未加 scope 前綴的保存值。
        scope, payload = "user", encoded
    encrypted = base64.b64decode(payload.encode("ascii"), validate=True)
    flags = 0x4 if scope == "machine" else 0
    return _crypt_with_windows_dpapi(encrypted, protect=False, flags=flags).decode("utf-8")


def now_text() -> str:
    return datetime.now().strftime("%H:%M:%S")


def split_board_values(text: str) -> list[str]:
    """看板支援中英文逗號、空白及換行，並保留輸入順序去重。"""
    values = re.split(r"[,，\s]+", text.strip())
    result: list[str] = []
    seen: set[str] = set()
    for value in values:
        value = value.strip()
        key = value.casefold()
        if value and key not in seen:
            seen.add(key)
            result.append(value)
    return result


def split_filter_values(text: str) -> list[str]:
    """型號及地區只以逗號/換行分隔，保留 iPhone 16 Pro 內部空白。"""
    values = re.split(r"[,，;；\n]+", text.strip())
    result: list[str] = []
    seen: set[str] = set()
    for value in values:
        value = value.strip()
        key = value.casefold()
        if value and key not in seen:
            seen.add(key)
            result.append(value)
    return result


def compact_text(text: str) -> str:
    """將型號轉成適合比對的格式，例如 iPhone 16 Pro -> iphone16pro。"""
    return re.sub(r"[^0-9a-z\u4e00-\u9fff]+", "", text.casefold())


def normalize_location(text: str) -> str:
    text = text.replace("臺", "台").strip()
    return re.sub(r"[縣市區鄉鎮]+$", "", text)


def safe_error(exc: Exception, secret: str = "") -> str:
    """避免例外訊息意外把 API Key 顯示在 Log。"""
    message = str(exc).replace("\n", " ").strip()
    if secret:
        message = message.replace(secret, "***")
    return message[:400] or exc.__class__.__name__


@dataclass
class Article:
    board: str
    title: str
    url: str
    author: str = ""
    list_date: str = ""
    published_at: str = ""
    content: str = ""


class SeenArticleStore:
    def __init__(self) -> None:
        app_dir = Path(QStandardPaths.writableLocation(QStandardPaths.AppDataLocation))
        self.path = app_dir / "seen_articles.json"
        self.data: dict[str, dict[str, str]] = {}
        self.load()

    @staticmethod
    def article_id(url: str) -> str:
        return Path(url).stem

    def load(self) -> None:
        try:
            if self.path.exists():
                raw = json.loads(self.path.read_text(encoding="utf-8"))
                self.data = raw if isinstance(raw, dict) else {}
        except (OSError, json.JSONDecodeError):
            # 紀錄檔損壞時不讓主程式無法啟動。
            self.data = {}

    def contains(self, url: str) -> bool:
        return self.article_id(url) in self.data

    def add(self, article: Article) -> None:
        self.data[self.article_id(article.url)] = {
            "url": article.url,
            "board": article.board,
            "title": article.title,
            "checked_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        }

    def save(self) -> None:
        self.path.parent.mkdir(parents=True, exist_ok=True)
        temp_path = self.path.with_suffix(".tmp")
        temp_path.write_text(
            json.dumps(self.data, ensure_ascii=False, indent=2), encoding="utf-8"
        )
        temp_path.replace(self.path)

    def clear(self) -> None:
        self.data = {}
        self.save()


class PTTBoardNotFound(RuntimeError):
    pass


class PTTMacShopCrawler:
    """PTT 多看板爬蟲；類別名稱沿用最初用途，但可搜尋任意公開看板。"""

    def __init__(self, delay_seconds: float = 0.45) -> None:
        self.delay_seconds = delay_seconds
        self.session = requests.Session()
        retry_strategy = Retry(
            total=2,
            backoff_factor=0.3,
            status_forcelist=[500, 502, 503, 504],
            raise_on_status=False,
        )
        adapter = HTTPAdapter(max_retries=retry_strategy)
        self.session.mount("https://", adapter)
        self.session.mount("http://", adapter)
        self.session.cookies.update({"over18": "1"})
        self.session.headers.update(
            {
                "User-Agent": (
                    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                    "AppleWebKit/537.36 Chrome/124 Safari/537.36 "
                    "PTT-Phone-Finder/1.0"
                )
            }
        )

    def _get(self, url: str) -> requests.Response:
        response = self.session.get(url, timeout=(8, 20))
        if response.status_code == 404:
            raise PTTBoardNotFound("看板或文章不存在（HTTP 404）")
        response.raise_for_status()
        return response

    def fetch_board_pages(
        self,
        board: str,
        page_count: int,
        stop_event: threading.Event,
        page_callback,
    ) -> list[Article]:
        url = f"{PTT_BASE}/bbs/{board}/index.html"
        articles: list[Article] = []
        known_urls: set[str] = set()

        for page_no in range(1, page_count + 1):
            if stop_event.is_set():
                break
            page_callback(page_no, page_count)
            response = self._get(url)
            soup = BeautifulSoup(response.text, "html.parser")

            # PTT 不存在的看板有時會回傳非預期頁面而非單純 404。
            if not soup.select("div.r-ent") and "不存在" in soup.get_text(" ", strip=True):
                raise PTTBoardNotFound("看板不存在")

            for entry in soup.select("div.r-ent"):
                anchor = entry.select_one("div.title a")
                if not anchor or not anchor.get("href"):
                    continue  # 已刪除文章沒有連結
                article_url = urljoin(PTT_BASE, anchor["href"])
                if article_url in known_urls:
                    continue
                known_urls.add(article_url)
                author_node = entry.select_one("div.author")
                date_node = entry.select_one("div.date")
                articles.append(
                    Article(
                        board=board,
                        title=anchor.get_text(" ", strip=True),
                        url=article_url,
                        author=author_node.get_text(strip=True) if author_node else "",
                        list_date=date_node.get_text(strip=True) if date_node else "",
                    )
                )

            previous = next(
                (
                    a
                    for a in soup.select("div.btn-group-paging a.btn")
                    if "上頁" in a.get_text(strip=True) and a.get("href")
                ),
                None,
            )
            if not previous:
                break
            url = urljoin(PTT_BASE, previous["href"])
            time.sleep(self.delay_seconds)
        return articles

    def fetch_article(self, article: Article) -> Article:
        response = self._get(article.url)
        soup = BeautifulSoup(response.text, "html.parser")
        main = soup.select_one("#main-content")
        if main is None:
            raise ValueError("找不到文章內容區塊，PTT HTML 結構可能已變更")

        meta_values = [node.get_text(" ", strip=True) for node in main.select(".article-meta-value")]
        if len(meta_values) >= 4:
            article.published_at = meta_values[3]

        # 移除推文、頁首欄位及簽名檔，減少送給 Gemini 的無關文字與 token。
        for node in main.select(".push, .article-metaline, .article-metaline-right"):
            node.decompose()
        raw_text = main.get_text("\n", strip=True)
        if "※ 發信站:" in raw_text:
            parts = raw_text.split("※ 發信站:", 1)
            body = parts[0]
            # 保留發信站之後作者可能追加的編輯/修文註記 (例如 ※ 編輯: ... (已售出))
            edits = re.findall(r"(※\s*編輯:[^\n]+)", parts[1])
            if edits:
                body = f"{body}\n" + "\n".join(edits)
            article.content = body.strip()[:18_000]
        else:
            article.content = raw_text[:18_000]
        time.sleep(self.delay_seconds)
        return article


def parse_published_datetime(value: str) -> datetime | None:
    """解析 PTT 文章頁的英文日期格式，解析不到時保留原始文字。"""
    text = value.strip()
    match = re.search(r"[A-Za-z]{3}\s+([A-Za-z]{3})\s+(\d{1,2})\s+(\d{2}:\d{2}:\d{2})\s+(\d{4})", text)
    if match:
        months = {
            "Jan": 1, "Feb": 2, "Mar": 3, "Apr": 4, "May": 5, "Jun": 6,
            "Jul": 7, "Aug": 8, "Sep": 9, "Oct": 10, "Nov": 11, "Dec": 12,
        }
        month = months.get(match.group(1).title())
        if month:
            return datetime.strptime(
                f"{match.group(4)}-{month:02d}-{int(match.group(2)):02d} {match.group(3)}",
                "%Y-%m-%d %H:%M:%S",
            )
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y/%m/%d %H:%M:%S"):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    return None


class PythonArticleClassifier:
    """使用本機規則從文章標題與內文抽取可核對的基本資訊。"""

    def __init__(self, keywords: list[str]) -> None:
        self.keywords = keywords

    @staticmethod
    def _unique(values: list[str]) -> list[str]:
        result: list[str] = []
        seen: set[str] = set()
        for value in values:
            key = value.casefold()
            if value and key not in seen:
                seen.add(key)
                result.append(value)
        return result

    @staticmethod
    def _extract_prices(text: str) -> list[int]:
        prices: list[int] = []
        for match in CHINESE_PRICE_PATTERN.finditer(text):
            tail = match.group(2) or "0"
            prices.append(int(match.group(1)) * 10_000 + int(tail))
        for match in PRICE_PATTERN.finditer(text):
            raw = match.group(1).replace(",", "")
            val = int(raw)
            # 排除年份（例如 2023年購入、2024/09 等非價格資訊）
            after = text[match.end():match.end() + 2]
            if after.startswith("年") or after.startswith("/") or after.startswith("-"):
                continue
            if 2018 <= val <= 2030 and not bool(re.search(r"(?:元|塊|新台幣|NT\$|TWD)", match.group(0), re.I)):
                continue

            before = text[max(0, match.start() - 10) : match.start()]
            explicit_currency = bool(re.search(r"(?:元|塊|新台幣|NT\$|TWD)", match.group(0), re.I))
            price_context = bool(re.search(r"(?:價格|售價|賣價|售出價|售|賣|收)", before))
            if (explicit_currency or price_context) and val not in prices:
                prices.append(val)
        return prices

    @staticmethod
    def _extract_capacities(text: str) -> list[str]:
        capacities = []
        for match in CAPACITY_PATTERN.finditer(text):
            num = match.group(1)
            unit = match.group(2).upper()
            if unit in ("T", "TB") or (num == "1024" and unit in ("G", "GB")):
                tb_val = "1" if num in ("1", "1024") else num
                capacities.append(f"{tb_val}TB")
            else:
                capacities.append(f"{num}GB")
        return PythonArticleClassifier._unique(capacities)

    @staticmethod
    def _extract_colors(text: str) -> list[str]:
        compact_value = compact_text(text)
        matched = [color for color in COLOR_TERMS_CHINESE if compact_text(color) in compact_value]
        # 英文顏色使用獨立單字邊界匹配，避免 "ordered" 誤判為 "red"
        for eng_color in COLOR_TERMS_ENGLISH:
            if re.search(rf"\b{re.escape(eng_color)}\b", text, re.I):
                matched.append(eng_color)
        matched = [
            color
            for color in matched
            if not any(
                color != other and compact_text(color) in compact_text(other)
                for other in matched
            )
        ]
        return PythonArticleClassifier._unique(matched)

    @staticmethod
    def _extract_sold_status(text: str) -> str:
        if any(term.casefold() in text.casefold() for term in UNSOLD_TERMS):
            return "未售出"
        if any(term.casefold() in text.casefold() for term in SOLD_TERMS):
            return "已售出"
        return "未判斷"

    def analyze(self, article: Article) -> dict[str, Any]:
        source = f"{article.title}\n{article.content}"
        matched_keywords = [
            keyword
            for keyword in self.keywords
            if compact_text(keyword) in compact_text(source)
        ]
        prices = self._extract_prices(source)
        capacities = self._extract_capacities(source)
        colors = self._extract_colors(source)
        sold_status = self._extract_sold_status(source)
        return {
            "matched_keywords": self._unique(matched_keywords),
            "model": "、".join(self._unique(matched_keywords)),
            "storage": "、".join(capacities),
            "price": prices[0] if prices else None,
            "prices_found": prices,
            "color": "、".join(colors),
            "sold_status": sold_status,
            "sold": sold_status == "已售出",
            "location": [],
            "condition": "Python 規則分類；雙擊列可開啟原文",
            "confidence": None,
        }


EXCEL_HEADERS = [
    "來源看板", "發文日期", "發文時間", "手機型號/命中關鍵字", "容量", "價格", "價格候選",
    "顏色", "售出狀態", "地區", "商品狀況", "信心度", "原始標題", "文章 URL", "文章內容",
]


def result_to_excel_row(data: dict[str, Any]) -> list[Any]:
    published_raw = str(data.get("published_at") or "")
    published_dt = parse_published_datetime(published_raw)
    if published_dt:
        published_date: Any = published_dt.date()
        published_time: Any = published_dt.time()
    else:
        published_date = data.get("list_date") or published_raw
        published_time = None
    locations = data.get("location")
    location_text = "、".join(str(value) for value in locations) if isinstance(locations, list) else str(locations or "")
    matched = data.get("matched_keywords")
    model_text = "、".join(str(value) for value in matched) if isinstance(matched, list) else str(data.get("model") or "")
    prices = data.get("prices_found")
    price_candidates = "、".join(f"{int(value):,}" for value in prices if isinstance(value, int)) if isinstance(prices, list) else ""
    sold_status = data.get("sold_status")
    if not sold_status and isinstance(data.get("sold"), bool):
        sold_status = "已售出" if data["sold"] else "未售出"
    confidence = data.get("confidence")
    return [
        data.get("board", ""),
        published_date,
        published_time,
        model_text,
        data.get("storage") or "",
        data.get("price") if isinstance(data.get("price"), int) and not isinstance(data.get("price"), bool) else None,
        price_candidates,
        data.get("color") or "",
        sold_status or "未判斷",
        location_text,
        data.get("condition") or "",
        confidence if isinstance(confidence, (int, float)) and not isinstance(confidence, bool) else None,
        data.get("title", ""),
        data.get("url", ""),
        data.get("content", ""),
    ]


def export_results_to_excel(path: str, results: list[dict[str, Any]]) -> None:
    """將搜尋結果輸出成可排序、可篩選的 Excel 工作表。"""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "搜尋結果"
    sheet.sheet_view.showGridLines = False
    sheet.append(EXCEL_HEADERS)
    for data in results:
        sheet.append(result_to_excel_row(data))

    header_fill = PatternFill("solid", fgColor="1F4E78")
    sold_fill = PatternFill("solid", fgColor="FCE4D6")
    header_font = Font(name="Aptos", bold=True, color="FFFFFF")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column in {11, 13, 15})
        row[1].number_format = "yyyy-mm-dd"
        row[2].number_format = "hh:mm:ss"
        row[5].number_format = "#,##0"
        if row[11].value is not None:
            row[11].number_format = "0%"
        if row[8].value == "已售出":
            row[8].fill = sold_fill
        if row[13].value:
            row[13].hyperlink = row[13].value
            row[13].style = "Hyperlink"

    widths = {1: 14, 2: 13, 3: 11, 4: 24, 5: 14, 6: 12, 7: 18, 8: 16, 9: 12, 10: 14, 11: 32, 12: 10, 13: 48, 14: 44, 15: 80}
    for column_index, width in widths.items():
        sheet.column_dimensions[get_column_letter(column_index)].width = width
    sheet.freeze_panes = "A2"
    if sheet.max_row >= 2:
        table = Table(displayName="PTTSearchResults", ref=f"A1:O{sheet.max_row}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
        sheet.add_table(table)
    else:
        sheet.auto_filter.ref = f"A1:O{sheet.max_row}"
    workbook.save(path)


class GeminiParser:
    RESPONSE_SCHEMA = {
        "type": "object",
        "properties": {
            "is_phone_sale": {"type": "boolean"},
            "model": {"type": ["string", "null"]},
            "storage": {"type": ["string", "null"]},
            "color": {"type": ["string", "null"]},
            "price": {"type": ["integer", "null"]},
            "location": {"type": "array", "items": {"type": "string"}},
            "condition": {"type": ["string", "null"]},
            "sold": {"type": "boolean"},
            "confidence": {"type": "number"},
        },
        "required": [
            "is_phone_sale",
            "model",
            "storage",
            "color",
            "price",
            "location",
            "condition",
            "sold",
            "confidence",
        ],
        "additionalProperties": False,
    }

    def __init__(self, api_key: str, model: str = GEMINI_MODEL) -> None:
        self.api_key = api_key
        self.model = model.strip() or GEMINI_MODEL
        self.client = genai.Client(
            api_key=api_key,
            http_options=types.HttpOptions(timeout=60_000),
        )

    def parse(self, article: Article) -> dict[str, Any]:
        prompt = f"""
你是台灣二手手機文章資料抽取器。只根據下方文章內容判斷，不可猜測。

規則：
1. 只有賣家販售手機本體時 is_phone_sale=true；徵求、交換、已售出、配件、維修、詢價或討論皆為 false。
2. price 必須是新台幣整數。27,500、27500、2萬7500 均轉成 27500；有多個價格時取手機本體實際售價；無法判斷為 null。
3. location 僅列文章明載的面交或交易縣市，使用簡稱（台北、新北、桃園、新竹、台中等）；多地點輸出多個元素，未寫則為空陣列。
4. 已售出或明確預訂完成時 sold=true。
5. model 輸出完整手機型號，storage 分開輸出容量，color 輸出顏色；未寫的欄位為 null。
6. condition 用繁體中文精簡摘要，最多 60 字，不加入文章沒有的資訊。
7. confidence 為 0 到 1。

來源看板：{article.board}
文章標題：{article.title}
文章內容：
{article.content}
""".strip()

        response = self.client.models.generate_content(
            model=self.model,
            contents=prompt,
            # 標準 JSON Schema（包含 nullable 欄位）要使用 response_json_schema。
            config={
                "response_mime_type": "application/json",
                "response_json_schema": self.RESPONSE_SCHEMA,
                "temperature": 0,
            },
        )
        raw = (response.text or "").strip()
        if raw.startswith("```"):
            raw = re.sub(r"^```(?:json)?\s*|\s*```$", "", raw, flags=re.I)
        result = json.loads(raw)
        if not isinstance(result, dict):
            raise ValueError("Gemini 回傳內容不是 JSON object")
        return result


@dataclass
class SearchOptions:
    api_key: str
    model: str
    mode: str
    boards: list[str]
    keywords: list[str]
    max_budget: int
    locations: list[str]
    pages: int
    unseen_only: bool


class SearchWorker(QThread):
    log = Signal(str)
    result_found = Signal(dict)
    summary = Signal(int, int)

    def __init__(self, options: SearchOptions) -> None:
        super().__init__()
        self.options = options
        self.stop_event = threading.Event()

    def request_stop(self) -> None:
        self.stop_event.set()

    def _log(self, message: str) -> None:
        self.log.emit(f"{now_text()} {message}")

    def _title_matches(self, title: str) -> bool:
        compact_title = compact_text(title)
        return any(compact_text(keyword) in compact_title for keyword in self.options.keywords)

    def _matching_keywords(self, text: str) -> list[str]:
        compact_value = compact_text(text)
        return [
            keyword
            for keyword in self.options.keywords
            if compact_text(keyword) in compact_value
        ]

    @staticmethod
    def _looks_like_sale(title: str) -> bool:
        compact = re.sub(r"\s+", "", title)
        # 僅排除標題已明確表達非販售或售完的文章，避免規則過嚴。
        blocked = ("[徵", "［徵", "[收購", "[交換", "[已售", "已售出")
        return not any(token.casefold() in compact.casefold() for token in blocked)

    def _model_matches(self, data: dict[str, Any]) -> bool:
        model = compact_text(str(data.get("model") or ""))
        storage = compact_text(str(data.get("storage") or ""))
        combined = model + storage
        return any(compact_text(keyword) in combined for keyword in self.options.keywords)

    def _location_matches(self, locations: list[Any]) -> bool:
        if not self.options.locations:
            return True
        targets = {normalize_location(value) for value in self.options.locations}
        actual = {normalize_location(str(value)) for value in locations}
        return bool(targets & actual)

    def _matches_final_conditions(self, data: dict[str, Any]) -> bool:
        if not data.get("is_phone_sale") or data.get("sold"):
            return False
        if not self._model_matches(data):
            return False
        price = data.get("price")
        if self.options.max_budget > 0:
            if not isinstance(price, int) or isinstance(price, bool):
                return False
            if price > self.options.max_budget:
                return False
        locations = data.get("location")
        if not isinstance(locations, list) or not self._location_matches(locations):
            return False
        return True

    def run(self) -> None:
        crawler = PTTMacShopCrawler()
        store = SeenArticleStore() if self.options.mode == SEARCH_MODE_GEMINI else None
        classifier = PythonArticleClassifier(self.options.keywords)
        parser: GeminiParser | None = None
        if self.options.mode == SEARCH_MODE_GEMINI:
            try:
                parser = GeminiParser(self.options.api_key, self.options.model)
            except Exception as exc:
                self._log(f"Gemini 初始化失敗：{safe_error(exc, self.options.api_key)}")
                self.summary.emit(0, 0)
                return

        candidate_count = 0
        match_count = 0
        if self.options.mode == SEARCH_MODE_PYTHON:
            self._log(f"開始搜尋，共 {len(self.options.boards)} 個看板；模式：純 Python 關鍵字搜尋")
        else:
            self._log(
                f"開始搜尋，共 {len(self.options.boards)} 個看板；Gemini 模型：{self.options.model}"
            )

        for board in self.options.boards:
            if self.stop_event.is_set():
                break
            self._log(f"[{board}] 正在驗證並讀取看板")
            try:
                articles = crawler.fetch_board_pages(
                    board,
                    self.options.pages,
                    self.stop_event,
                    lambda current, total, b=board: self._log(
                        f"[{b}] 正在讀取第 {current} / {total} 頁"
                    ),
                )
                self._log(f"[{board}] 取得 {len(articles)} 篇文章標題")
            except PTTBoardNotFound as exc:
                self._log(f"找不到看板 {board}，已略過：{safe_error(exc)}")
                continue
            except requests.RequestException as exc:
                self._log(f"[{board}] PTT 連線失敗，已略過：{safe_error(exc)}")
                continue
            except Exception as exc:
                self._log(f"[{board}] 讀取失敗，已略過：{safe_error(exc)}")
                continue

            candidates = [
                article
                for article in articles
                if self._title_matches(article.title)
                and self._looks_like_sale(article.title)
                and not (
                    self.options.mode == SEARCH_MODE_GEMINI
                    and self.options.unseen_only
                    and store is not None
                    and store.contains(article.url)
                )
            ]
            candidate_count += len(candidates)
            self._log(f"[{board}] 初步篩選出 {len(candidates)} 篇可能相關文章")

            for index, article in enumerate(candidates, start=1):
                if self.stop_event.is_set():
                    break
                if self.options.mode == SEARCH_MODE_PYTHON:
                    try:
                        if self.stop_event.is_set():
                            break
                        article = crawler.fetch_article(article)
                        data = classifier.analyze(article)
                        price = data.get("price")
                        if self.options.max_budget > 0 and isinstance(price, int) and price > self.options.max_budget:
                            continue
                        result = asdict(article)
                        result.update(data)
                        match_count += 1
                        self.result_found.emit(result)
                        self._log(f"[{board}] 找到關鍵字商品：{article.title}")
                    except requests.RequestException as exc:
                        self._log(f"[{board}] 文章下載失敗，繼續下一篇：{safe_error(exc)}")
                    except Exception as exc:
                        self._log(f"[{board}] Python 分類失敗，繼續下一篇：{safe_error(exc)}")
                    continue

                self._log(f"[{board}] Gemini 分析第 {index} / {len(candidates)} 篇")
                try:
                    article = crawler.fetch_article(article)
                    if not article.content:
                        raise ValueError("文章沒有可分析內容")
                    if parser is None:
                        raise RuntimeError("Gemini parser 尚未初始化")
                    data = parser.parse(article)
                    # 成功完成 AI 分析後才記為已讀；暫時性 API 錯誤可在下次重試。
                    if store is None:
                        raise RuntimeError("Gemini seen store 尚未初始化")
                    store.add(article)
                    store.save()
                    if self._matches_final_conditions(data):
                        match_count += 1
                        result = asdict(article)
                        result.update(data)
                        self.result_found.emit(result)
                        self._log(f"[{board}] 發現符合條件商品：{article.title}")
                except requests.RequestException as exc:
                    self._log(f"[{board}] 文章下載失敗，繼續下一篇：{safe_error(exc)}")
                except (json.JSONDecodeError, ValueError) as exc:
                    self._log(f"[{board}] Gemini JSON 解析失敗，繼續下一篇：{safe_error(exc)}")
                except Exception as exc:
                    self._log(
                        f"[{board}] Gemini 分析失敗，繼續下一篇："
                        f"{safe_error(exc, self.options.api_key)}"
                    )

        if self.stop_event.is_set():
            self._log(f"搜尋已停止，目前找到 {match_count} 筆符合商品")
        else:
            self._log(f"搜尋完成，共找到 {match_count} 筆符合條件商品")
        self.summary.emit(match_count, candidate_count)


class NumericTableWidgetItem(QTableWidgetItem):
    """讓價格欄依數值排序，而不是依字串排序。"""

    def __lt__(self, other: QTableWidgetItem) -> bool:
        left = self.data(Qt.ItemDataRole.UserRole)
        right = other.data(Qt.ItemDataRole.UserRole)
        if isinstance(left, (int, float)) and isinstance(right, (int, float)):
            return left < right
        return super().__lt__(other)


class MainWindow(QMainWindow):
    COLUMNS = [
        "來源看板", "發文日期", "發文時間", "手機型號", "容量", "價格", "價格候選",
        "顏色", "售出狀態", "地區", "商品狀況", "信心度", "原始標題", "文章 URL",
    ]

    def __init__(self) -> None:
        super().__init__()
        self.worker: SearchWorker | None = None
        self.results_data: list[dict[str, Any]] = []
        # 明確使用使用者 AppData 的 INI，避免受限 Windows 帳號無法寫入 Registry。
        self.settings = QSettings(
            QSettings.Format.IniFormat,
            QSettings.Scope.UserScope,
            "LocalTools",
            "PTTPhoneFinder",
        )
        self.setWindowTitle("PTT 二手手機搜尋工具")
        self.resize(1280, 820)
        self._build_ui()

    def _build_ui(self) -> None:
        root = QWidget()
        self.setCentralWidget(root)
        root_layout = QVBoxLayout(root)

        settings = QGroupBox("搜尋設定")
        settings_layout = QGridLayout(settings)

        self.mode_combo = QComboBox()
        for label, mode in SEARCH_MODES:
            self.mode_combo.addItem(label, mode)
        saved_mode = str(self.settings.value("search_mode", SEARCH_MODE_PYTHON) or SEARCH_MODE_PYTHON)
        mode_index = self.mode_combo.findData(saved_mode)
        self.mode_combo.setCurrentIndex(mode_index if mode_index >= 0 else 0)
        settings_layout.addWidget(QLabel("搜尋模式"), 0, 0)
        settings_layout.addWidget(self.mode_combo, 0, 1, 1, 3)

        self.api_key_edit = QLineEdit()
        self.api_key_edit.setEchoMode(QLineEdit.EchoMode.Password)
        saved_api_key = str(self.settings.value("gemini_api_key_dpapi", "") or "")
        if saved_api_key:
            try:
                self.api_key_edit.setText(unprotect_api_key(saved_api_key))
            except Exception:
                # 舊 Windows 使用者或資料損壞時，忽略無法解密的值，不把例外帶到 GUI。
                self.settings.remove("gemini_api_key_dpapi")
        self.api_key_edit.setPlaceholderText("貼上 Gemini API Key（不會顯示在 Log）")
        self.show_key = QCheckBox("顯示 Key")
        self.remember_key = QCheckBox("記住 Key（Windows 加密儲存）")
        self.remember_key.setChecked(True)
        self.show_key.toggled.connect(
            lambda checked: self.api_key_edit.setEchoMode(
                QLineEdit.EchoMode.Normal if checked else QLineEdit.EchoMode.Password
            )
        )
        key_row = QHBoxLayout()
        key_row.addWidget(self.api_key_edit)
        key_row.addWidget(self.show_key)
        key_row.addWidget(self.remember_key)
        settings_layout.addWidget(QLabel("Gemini API Key（AI 模式）"), 1, 0)
        settings_layout.addLayout(key_row, 1, 1, 1, 3)

        self.model_combo = QComboBox()
        self.model_combo.setEditable(True)
        self.model_combo.setInsertPolicy(QComboBox.InsertPolicy.NoInsert)
        self.model_combo.setToolTip("可選擇或直接輸入 Google AI Studio 可用的模型代號")
        for label, model in GEMINI_MODELS:
            self.model_combo.addItem(label, model)
        saved_model = str(self.settings.value("gemini_model", GEMINI_MODEL) or GEMINI_MODEL)
        model_index = self.model_combo.findData(saved_model)
        if model_index >= 0:
            self.model_combo.setCurrentIndex(model_index)
        else:
            self.model_combo.setEditText(saved_model)
        settings_layout.addWidget(QLabel("Gemini 模型（AI 模式）"), 2, 0)
        settings_layout.addWidget(self.model_combo, 2, 1, 1, 3)

        board_box = QGroupBox("常用看板（可複選）")
        board_layout = QHBoxLayout(board_box)
        self.board_checks: dict[str, QCheckBox] = {}
        for board in DEFAULT_BOARDS:
            check = QCheckBox(board)
            check.setChecked(board in {"MacShop", "mobilesales"})
            self.board_checks[board] = check
            board_layout.addWidget(check)
        board_layout.addStretch()
        settings_layout.addWidget(board_box, 3, 0, 1, 4)

        self.custom_boards_edit = QLineEdit()
        self.custom_boards_edit.setPlaceholderText("例如：MacShop, mobilesales（可用逗號、空白或換行分隔）")
        settings_layout.addWidget(QLabel("自訂看板"), 4, 0)
        settings_layout.addWidget(self.custom_boards_edit, 4, 1, 1, 3)

        self.keyword_edit = QLineEdit("iPhone 16 Pro")
        self.keyword_edit.setPlaceholderText("多個型號可用逗號分隔")
        self.budget_spin = QSpinBox()
        self.budget_spin.setRange(0, 10_000_000)
        self.budget_spin.setValue(30_000)
        self.budget_spin.setSingleStep(1_000)
        self.budget_spin.setSpecialValueText("不限")
        self.budget_spin.setSuffix(" 元")
        self.location_edit = QLineEdit("台北, 新竹")
        self.location_edit.setPlaceholderText("留空代表不限地區")
        self.pages_spin = QSpinBox()
        self.pages_spin.setRange(1, 20)
        self.pages_spin.setValue(3)
        self.pages_spin.setSuffix(" 頁 / 看板")

        settings_layout.addWidget(QLabel("手機型號 / 關鍵字"), 5, 0)
        settings_layout.addWidget(self.keyword_edit, 5, 1)
        settings_layout.addWidget(QLabel("最高預算（AI 模式）"), 5, 2)
        settings_layout.addWidget(self.budget_spin, 5, 3)
        settings_layout.addWidget(QLabel("限定地區（AI 模式）"), 6, 0)
        settings_layout.addWidget(self.location_edit, 6, 1)
        settings_layout.addWidget(QLabel("搜尋最新頁數"), 6, 2)
        settings_layout.addWidget(self.pages_spin, 6, 3)

        self.unseen_only_check = QCheckBox("只分析尚未看過的文章（僅 Gemini 模式）")
        self.unseen_only_check.setChecked(True)
        settings_layout.addWidget(self.unseen_only_check, 7, 0, 1, 2)

        button_row = QHBoxLayout()
        self.start_button = QPushButton("開始搜尋")
        self.stop_button = QPushButton("停止搜尋")
        self.clear_button = QPushButton("清除結果")
        self.clear_seen_button = QPushButton("清除已讀紀錄")
        self.export_button = QPushButton("匯出 Excel")
        self.stop_button.setEnabled(False)
        button_row.addWidget(self.start_button)
        button_row.addWidget(self.stop_button)
        button_row.addWidget(self.clear_button)
        button_row.addWidget(self.clear_seen_button)
        button_row.addWidget(self.export_button)
        button_row.addStretch()
        settings_layout.addLayout(button_row, 7, 2, 1, 2)
        root_layout.addWidget(settings)

        self.table = QTableWidget(0, len(self.COLUMNS))
        self.table.setHorizontalHeaderLabels(self.COLUMNS)
        self.table.setSortingEnabled(True)
        self.table.setSelectionBehavior(QTableWidget.SelectionBehavior.SelectRows)
        self.table.setEditTriggers(QTableWidget.EditTrigger.NoEditTriggers)
        self.table.verticalHeader().setVisible(False)
        header = self.table.horizontalHeader()
        header.setSectionResizeMode(QHeaderView.ResizeMode.ResizeToContents)
        header.setSectionResizeMode(10, QHeaderView.ResizeMode.Stretch)
        header.setSectionResizeMode(12, QHeaderView.ResizeMode.Stretch)
        self.table.setColumnHidden(13, True)
        self.table.cellDoubleClicked.connect(self.open_article)

        self.log_list = QListWidget()
        splitter = QSplitter(Qt.Orientation.Vertical)
        splitter.addWidget(self.table)
        splitter.addWidget(self.log_list)
        splitter.setSizes([500, 180])
        root_layout.addWidget(splitter, 1)

        self.statusBar().showMessage("就緒；雙擊搜尋結果可開啟 PTT 原文")
        self.start_button.clicked.connect(self.start_search)
        self.stop_button.clicked.connect(self.stop_search)
        self.clear_button.clicked.connect(self.clear_results)
        self.clear_seen_button.clicked.connect(self.clear_seen)
        self.export_button.clicked.connect(self.export_excel)
        self.mode_combo.currentIndexChanged.connect(self._on_mode_changed)
        self._on_mode_changed()

    def selected_boards(self) -> list[str]:
        selected = [name for name, checkbox in self.board_checks.items() if checkbox.isChecked()]
        selected.extend(split_board_values(self.custom_boards_edit.text()))
        result: list[str] = []
        seen: set[str] = set()
        for board in selected:
            # PTT 看板名稱只接受安全的常見字元，避免產生異常 URL。
            if not re.fullmatch(r"[A-Za-z0-9_.-]+", board):
                continue
            key = board.casefold()
            if key not in seen:
                seen.add(key)
                result.append(board)
        return result

    def selected_model(self) -> str:
        current_index = self.model_combo.currentIndex()
        current_text = self.model_combo.currentText().strip()
        if current_index >= 0 and current_text == self.model_combo.itemText(current_index):
            return str(self.model_combo.itemData(current_index) or "").strip()
        return current_text

    def selected_mode(self) -> str:
        return str(self.mode_combo.currentData() or SEARCH_MODE_PYTHON)

    def _on_mode_changed(self) -> None:
        ai_enabled = self.selected_mode() == SEARCH_MODE_GEMINI
        for widget in (self.api_key_edit, self.show_key, self.remember_key, self.model_combo):
            widget.setEnabled(ai_enabled)
        self.budget_spin.setEnabled(ai_enabled)
        self.location_edit.setEnabled(ai_enabled)
        self.unseen_only_check.setEnabled(ai_enabled)
        self.clear_seen_button.setEnabled(ai_enabled)
        if hasattr(self, "table"):
            self.table.setHorizontalHeaderItem(
                3,
                QTableWidgetItem("手機型號" if ai_enabled else "命中關鍵字"),
            )
        if ai_enabled:
            self.statusBar().showMessage("AI 模式：需要 Gemini API Key；可選擇模型")
        else:
            self.statusBar().showMessage("純 Python 模式：不需要 Gemini API Key，直接列出關鍵字文章")

    def persist_preferences(self) -> bool:
        """保存模型與使用者選擇；API Key 僅以 DPAPI 加密後保存。"""
        self.settings.setValue("search_mode", self.selected_mode())
        model = self.selected_model()
        if model:
            self.settings.setValue("gemini_model", model)

        api_key = self.api_key_edit.text().strip()
        self.settings.remove("gemini_api_key")
        if not self.remember_key.isChecked() or not api_key:
            self.settings.remove("gemini_api_key_dpapi")
            self.settings.sync()
            return True

        try:
            self.settings.setValue("gemini_api_key_dpapi", protect_api_key(api_key))
            self.settings.sync()
            return True
        except Exception as exc:
            self.settings.remove("gemini_api_key_dpapi")
            self.settings.sync()
            QMessageBox.warning(
                self,
                "API Key 保存失敗",
                f"Windows 加密儲存失敗，本次仍可搜尋：{safe_error(exc)}",
            )
            return False

    def append_log(self, text: str) -> None:
        self.log_list.addItem(text)
        self.log_list.scrollToBottom()
        self.statusBar().showMessage(text)

    def start_search(self) -> None:
        if self.worker and self.worker.isRunning():
            return
        mode = self.selected_mode()
        api_key = self.api_key_edit.text().strip()
        model = self.selected_model()
        boards = self.selected_boards()
        keywords = split_filter_values(self.keyword_edit.text())
        if mode == SEARCH_MODE_GEMINI and not api_key:
            QMessageBox.warning(self, "缺少設定", "請先輸入 Gemini API Key。")
            return
        if mode == SEARCH_MODE_GEMINI and (
            not model or not re.fullmatch(r"[A-Za-z0-9_.:-]+", model)
        ):
            QMessageBox.warning(self, "模型設定錯誤", "請選擇或輸入有效的 Gemini 模型代號。")
            return
        if not boards:
            QMessageBox.warning(self, "缺少設定", "請至少選擇或輸入一個有效的 PTT 看板。")
            return
        if not keywords:
            QMessageBox.warning(self, "缺少設定", "請至少輸入一個手機型號或關鍵字。")
            return

        self.persist_preferences()
        options = SearchOptions(
            api_key=api_key,
            model=model,
            mode=mode,
            boards=boards,
            keywords=keywords,
            max_budget=self.budget_spin.value(),
            locations=split_filter_values(self.location_edit.text()),
            pages=self.pages_spin.value(),
            unseen_only=self.unseen_only_check.isChecked(),
        )
        self.worker = SearchWorker(options)
        self.worker.log.connect(self.append_log)
        self.worker.result_found.connect(self.add_result)
        self.worker.summary.connect(self.show_summary)
        self.worker.finished.connect(self.worker_finished)
        self.start_button.setEnabled(False)
        self.stop_button.setEnabled(True)
        self.worker.start()

    def stop_search(self) -> None:
        if self.worker and self.worker.isRunning():
            self.worker.request_stop()
            self.stop_button.setEnabled(False)
            self.append_log(f"{now_text()} 已要求安全停止；正在等待目前的網路請求結束")

    def worker_finished(self) -> None:
        self.start_button.setEnabled(True)
        self.stop_button.setEnabled(False)
        self.worker = None

    def show_summary(self, matches: int, candidates: int) -> None:
        self.statusBar().showMessage(f"完成：{candidates} 篇候選文章，{matches} 筆符合條件")

    def clear_results(self) -> None:
        self.results_data.clear()
        self.table.setRowCount(0)

    def add_result(self, data: dict[str, Any]) -> None:
        self.results_data.append(data)
        self.table.setSortingEnabled(False)
        row = self.table.rowCount()
        self.table.insertRow(row)
        price = data.get("price")
        published_raw = str(data.get("published_at") or "")
        published_dt = parse_published_datetime(published_raw)
        published_date = published_dt.strftime("%Y-%m-%d") if published_dt else str(data.get("list_date") or published_raw)
        published_time = published_dt.strftime("%H:%M:%S") if published_dt else ""
        locations = "、".join(str(value) for value in data.get("location", [])) or "未提供"
        matched = data.get("matched_keywords")
        model_text = "、".join(str(value) for value in matched) if isinstance(matched, list) else str(data.get("model") or "")
        prices = data.get("prices_found")
        price_candidates = "、".join(f"{int(value):,}" for value in prices if isinstance(value, int)) if isinstance(prices, list) else ""
        sold_status = data.get("sold_status")
        if not sold_status and isinstance(data.get("sold"), bool):
            sold_status = "已售出" if data["sold"] else "未售出"
        confidence = data.get("confidence")
        values = [
            data.get("board", ""),
            published_date,
            published_time,
            model_text or "未知",
            data.get("storage") or "未知",
            f"{price:,}" if isinstance(price, int) else "價格未知",
            price_candidates,
            data.get("color") or "未提供",
            sold_status or "未判斷",
            locations,
            data.get("condition") or "未提供",
            f"{float(confidence):.0%}" if isinstance(confidence, (int, float)) else "",
            data.get("title", ""),
            data.get("url", ""),
        ]
        for column, value in enumerate(values):
            item = NumericTableWidgetItem(str(value)) if column == 5 else QTableWidgetItem(str(value))
            if column == 5 and isinstance(price, int):
                item.setData(Qt.ItemDataRole.UserRole, price)
                item.setTextAlignment(Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter)
            if column == 13:
                item.setData(Qt.ItemDataRole.UserRole, data.get("url", ""))
            self.table.setItem(row, column, item)
        self.table.setSortingEnabled(True)

    def open_article(self, row: int, _column: int) -> None:
        item = self.table.item(row, 13)
        if item:
            url = item.data(Qt.ItemDataRole.UserRole) or item.text()
            if url:
                QDesktopServices.openUrl(QUrl(str(url)))

    def export_excel(self) -> None:
        if not self.results_data:
            QMessageBox.information(self, "沒有搜尋結果", "請先執行搜尋，再匯出 Excel。")
            return
        default_name = f"PTT手機搜尋_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
        path, _ = QFileDialog.getSaveFileName(
            self,
            "匯出搜尋結果",
            str(Path.home() / "Downloads" / default_name),
            "Excel 活頁簿 (*.xlsx)",
        )
        if not path:
            return
        if not path.lower().endswith(".xlsx"):
            path += ".xlsx"
        try:
            export_results_to_excel(path, self.results_data)
            self.append_log(f"{now_text()} 已匯出 Excel：{path}")
            QMessageBox.information(self, "匯出完成", f"已保存 Excel 檔案：\n{path}")
        except Exception as exc:
            QMessageBox.critical(self, "匯出失敗", f"無法建立 Excel：{safe_error(exc)}")

    def clear_seen(self) -> None:
        answer = QMessageBox.question(
            self,
            "確認清除",
            "確定要清除所有已讀文章紀錄嗎？清除後可能再次呼叫 Gemini 分析舊文章。",
        )
        if answer == QMessageBox.StandardButton.Yes:
            try:
                store = SeenArticleStore()
                store.clear()
                self.append_log(f"{now_text()} 已清除已讀文章紀錄")
            except OSError as exc:
                QMessageBox.critical(self, "清除失敗", safe_error(exc))

    def closeEvent(self, event) -> None:  # type: ignore[override]
        if self.worker and self.worker.isRunning():
            answer = QMessageBox.question(
                self,
                "搜尋仍在進行",
                "搜尋仍在進行，是否要求停止並關閉？",
            )
            if answer != QMessageBox.StandardButton.Yes:
                event.ignore()
                return
            self.worker.request_stop()
            # 不使用 terminate；讓 Worker 有機會安全離開。網路 timeout 最長約 60 秒。
            if not self.worker.wait(2500):
                self.worker.finished.connect(QApplication.instance().quit)
                self.hide()
                event.ignore()
                return
        self.persist_preferences()
        event.accept()


def main() -> None:
    app = QApplication(sys.argv)
    app.setApplicationName("PTTPhoneFinder")
    app.setOrganizationName("LocalTools")
    window = MainWindow()
    window.show()
    sys.exit(app.exec())


if __name__ == "__main__":
    main()
